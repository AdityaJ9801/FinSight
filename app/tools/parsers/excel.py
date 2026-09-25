"""Excel parsing via openpyxl directly (not pandas) so we keep cell coordinates for
provenance and formatting hints (indentation, bold) for subtotal detection (design doc
§6.4: "openpyxl (merged cells, indentation level, bold = subtotal hint)").

Processes EVERY sheet in the workbook, not just the first -- a single workbook commonly
carries P&L/BS/CF as separate sheets, each needing its own header-row detection since
metadata blocks (and their length) vary sheet to sheet.
"""
from __future__ import annotations

import openpyxl

from app.tools.parsers import RawRow, RawTable
from app.tools.registry import tool
from app.tools.table_detect import detect_header_row_heuristic


def _value_grid(rows) -> list[list]:
    return [[c.value for c in row] for row in rows]


def read_excel_grid(file_path: str, max_rows: int = 30) -> dict[str, list[list]]:
    """Raw cell-value preview per sheet, no header assumed -- used for header-row
    detection (heuristic or LLM) before the real parse."""
    wb = openpyxl.load_workbook(file_path, data_only=True, read_only=True)
    try:
        return {ws.title: _value_grid(list(ws.iter_rows(max_row=max_rows))) for ws in wb.worksheets}
    finally:
        wb.close()


def _build_sheet_table(rows, header_row_idx: int, sheet_name: str) -> RawTable | None:
    if header_row_idx < 0 or header_row_idx >= len(rows):
        header_row_idx = 0
    header_row = rows[header_row_idx]
    periods = [str(c.value).strip() for c in header_row[1:] if c.value not in (None, "")]
    if not periods:
        return None

    raw_rows: list[RawRow] = []
    for row in rows[header_row_idx + 1:]:
        label_cell = row[0]
        if label_cell.value in (None, ""):
            continue
        label = str(label_cell.value).strip()

        values: dict[str, float] = {}
        for j, period in enumerate(periods):
            if j + 1 >= len(row):
                break
            cell = row[j + 1]
            if cell.value is None:
                continue
            try:
                values[period] = float(cell.value)
            except (TypeError, ValueError):
                continue
        if not values:
            continue  # note/footnote-only row, not real data (see csv_tool's equivalent filter)

        indent = 0
        try:
            indent = int(label_cell.alignment.indent or 0) if label_cell.alignment else 0
        except (TypeError, ValueError):
            indent = 0
        is_bold = bool(label_cell.font and label_cell.font.bold)

        raw_rows.append(RawRow(
            row_idx=label_cell.row, label=label, values=values,
            is_subtotal=is_bold, indent_level=indent,
            source_ref={"sheet": sheet_name, "cell": label_cell.coordinate},
        ))

    return RawTable(rows=raw_rows, periods=periods, name=sheet_name) if raw_rows else None


def read_excel(file_path: str, header_overrides: dict[str, int] | None = None) -> dict[str, dict]:
    """Returns {sheet_name: {"table": RawTable, "header_row_idx": int, "confident": bool}}
    for every sheet that yields usable data. `confident=False` means the deterministic
    heuristic couldn't find the header row and row 0 was used as a fallback -- the
    extractor may re-call with `header_overrides={sheet_name: llm_chosen_idx}` to correct
    a specific sheet using an LLM-picked offset instead.
    """
    header_overrides = header_overrides or {}
    wb = openpyxl.load_workbook(file_path, data_only=True)
    try:
        results: dict[str, dict] = {}
        for ws in wb.worksheets:
            rows = list(ws.iter_rows())
            if len(rows) < 2:
                continue

            if ws.title in header_overrides:
                header_idx, confident = header_overrides[ws.title], True
            else:
                detected = detect_header_row_heuristic(_value_grid(rows))
                header_idx, confident = (detected, True) if detected is not None else (0, False)

            table = _build_sheet_table(rows, header_idx, ws.title)
            if table is not None:
                results[ws.title] = {"table": table, "header_row_idx": header_idx, "confident": confident}
        return results
    finally:
        wb.close()


tool("excel.read_sheets", allowed_agents=["extractor"])(read_excel)
tool("excel.read_grid", allowed_agents=["extractor"])(read_excel_grid)
